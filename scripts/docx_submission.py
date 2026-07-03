"""Build submission-ready DOCX variants (full + shortened for plagiarism check)."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from kkp_docx_fill import _find_paragraph_index, _normalize
from kkp_meta import DEFAULT_META, KkpMeta


def submission_stem(meta: KkpMeta = DEFAULT_META, *, year: str | None = None) -> str:
    """Official folder/file stem, e.g. 2026_Б_ПІ_ККП_ПЗПІ_23_5_Ткач_Л_Я."""
    yr = year or meta.year
    group = meta.group.replace("-", "_")
    initials = meta.student_short.replace(".", "").replace(" ", "_")
    return f"{yr}_Б_ПІ_ККП_{group}_{initials}"


def submission_paths(docs_dir: Path, meta: KkpMeta = DEFAULT_META) -> tuple[Path, Path]:
    stem = submission_stem(meta)
    return (
        docs_dir / f"{stem}.docx",
        docs_dir / f"{stem}_скорочений.docx",
    )


def _find_body_child_index(doc: Document, prefix: str, *, start: int = 0) -> int | None:
    needle = _normalize(prefix)
    body = doc.element.body
    for index, child in enumerate(body):
        if index < start:
            continue
        if not child.tag.endswith("}p"):
            continue
        text = _normalize(Paragraph(child, doc).text)
        if text.startswith(needle):
            return index
    return None


def _delete_body_children(doc: Document, start: int, end: int | None = None) -> None:
    body = doc.element.body
    end_index = len(body) if end is None else end
    for child in list(body)[start:end_index][::-1]:
        body.remove(child)


def _sect_pr_from_paragraph(paragraph: Paragraph):
    p_pr = paragraph._element.pPr
    if p_pr is None:
        return None
    sect = p_pr.find(qn("w:sectPr"))
    return deepcopy(sect) if sect is not None else None


def _extract_section_breaks(doc: Document) -> tuple[object | None, object | None]:
    """Return (section break before abstract, final document section properties)."""
    abstract_idx = _find_body_child_index(doc, "РЕФЕРАТ")
    title_break = None
    if abstract_idx is not None and abstract_idx > 0:
        prev = doc.element.body[abstract_idx - 1]
        if prev.tag.endswith("}p"):
            title_break = _sect_pr_from_paragraph(Paragraph(prev, doc))

    final_sect = None
    for child in reversed(doc.element.body):
        if child.tag == qn("w:sectPr"):
            final_sect = deepcopy(child)
            break
    return title_break, final_sect


def _apply_section_break(paragraph: Paragraph, sect_pr) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    existing = p_pr.find(qn("w:sectPr"))
    if existing is not None:
        p_pr.remove(existing)
    page_break = p_pr.find(qn("w:pageBreakBefore"))
    if page_break is not None:
        p_pr.remove(page_break)
    p_pr.append(deepcopy(sect_pr))


def _ensure_final_section_properties(doc: Document, final_sect) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            body.remove(child)
    if final_sect is not None:
        body.append(deepcopy(final_sect))


def _restore_title_page_section_break(doc: Document, title_break, final_sect) -> None:
    """Re-insert the section break that lived before РЕФЕРАТ in the full report."""
    if title_break is None:
        return
    year_idx = _find_paragraph_index(doc, f"{DEFAULT_META.year} р.")
    if year_idx is None:
        year_idx = _find_paragraph_index(doc, "2026 р.")
    if year_idx is None:
        return
    _apply_section_break(doc.paragraphs[year_idx], title_break)

    abstract_idx = _find_paragraph_index(doc, "РЕФЕРАТ")
    if abstract_idx is not None:
        abstract_ppr = doc.paragraphs[abstract_idx]._element.pPr
        if abstract_ppr is not None:
            page_break = abstract_ppr.find(qn("w:pageBreakBefore"))
            if page_break is not None:
                abstract_ppr.remove(page_break)

    _ensure_final_section_properties(doc, final_sect)


def build_shortened_report(src: Path, dst: Path) -> Path:
    """Drop sections excluded from plagiarism check (TOC, abbreviations, refs, appendices)."""
    doc = Document(str(src))
    title_break, final_sect = _extract_section_breaks(doc)

    # Bibliography and appendices (from heading to end; final sectPr is restored later).
    refs = _find_body_child_index(doc, "ПЕРЕЛІК ДЖЕРЕЛ ПОСИЛАННЯ")
    if refs is not None:
        _delete_body_children(doc, refs)

    appendices = _find_body_child_index(doc, "ДОДАТ")
    if appendices is not None:
        _delete_body_children(doc, appendices)

    # Abbreviations section (heading), not the TOC line "Перелік скорочень\t8".
    abbrev = _find_body_child_index(doc, "ПЕРЕЛІК СКОРОЧЕНЬ")
    intro = _find_body_child_index(doc, "ВСТУП")
    if abbrev is not None and intro is not None and abbrev < intro:
        _delete_body_children(doc, abbrev, intro)

    # Table of contents block.
    toc = _find_body_child_index(doc, "ЗМІСТ")
    if toc is not None:
        next_section = _find_body_child_index(doc, "ПЕРЕЛІК СКОРОЧЕНЬ", start=toc + 1)
        if next_section is None:
            next_section = _find_body_child_index(doc, "ВСТУП", start=toc + 1)
        if next_section is not None:
            _delete_body_children(doc, toc, next_section)

    # Assignment sheet (second title block through signatures before abstract).
    abstract = _find_body_child_index(doc, "РЕФЕРАТ")
    if abstract is not None:
        year_idx = _find_body_child_index(doc, "2026 р.")
        if year_idx is None:
            year_idx = _find_body_child_index(doc, f"{DEFAULT_META.year} р.")
        if year_idx is not None:
            _delete_body_children(doc, year_idx + 1, abstract)

    _restore_title_page_section_break(doc, title_break, final_sect)

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return dst


def prepare_submission_docs(
    source: Path,
    docs_dir: Path,
    *,
    meta: KkpMeta = DEFAULT_META,
) -> tuple[Path, Path]:
    full_path, short_path = submission_paths(docs_dir, meta)
    shutil.copy2(source, full_path)
    build_shortened_report(source, short_path)
    return full_path, short_path


def main() -> None:
    import argparse

    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description="Prepare full and shortened submission DOCX files."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=root / "docs" / "2026_Б_ККП_ПЗПІ-23-5_Ткач_Л_Я.docx",
        help="Full report DOCX to copy and shorten",
    )
    parser.add_argument(
        "--docs-dir",
        type=Path,
        default=root / "docs",
        help="Output directory",
    )
    args = parser.parse_args()
    full_path, short_path = prepare_submission_docs(args.source.resolve(), args.docs_dir.resolve())
    print(f"Full:        {full_path}")
    print(f"Shortened:   {short_path}")


if __name__ == "__main__":
    main()
