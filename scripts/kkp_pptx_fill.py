"""Fill official KKP presentation template with project-specific data."""

from __future__ import annotations

import struct
from pathlib import Path

from kkp_meta import KkpMeta
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Slides that show a figure beside the body text (two-column layout).
IMAGE_SLIDES: dict[int, str] = {
    5: "pipeline.png",
    8: "metrics_comparison.png",
    9: "gradcam_examples.png",
    10: "demo_overview.png",
}

# Inches — tuned for 10.00" × 5.62" template; title stays full width.
_TEXT_LEFT = 0.34
_TEXT_TOP = 1.12
_TEXT_WIDTH_FULL = 9.32
_TEXT_WIDTH_SPLIT = 4.55
_TEXT_HEIGHT = 4.2
_IMAGE_LEFT = 5.05
_IMAGE_TOP = 1.12
_IMAGE_MAX_WIDTH = 4.55
_IMAGE_MAX_HEIGHT = 4.2
_BODY_FONT_SPLIT = 13
_BODY_FONT_FULL = 15


def _set_shape_text(
    shape,
    text: str,
    *,
    font_size: int = 18,
    bold: bool = False,
    color: RGBColor | None = None,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            if color is not None:
                run.font.color.rgb = color


def _layout_title_slide(title_shape, body_shape) -> None:
    """Match KKP template: centered topic and author block."""
    title_shape.left = Inches(0.5)
    title_shape.top = Inches(0.85)
    title_shape.width = Inches(9.0)
    title_shape.height = Inches(1.0)
    body_shape.left = Inches(2.2)
    body_shape.top = Inches(3.55)
    body_shape.width = Inches(5.6)
    body_shape.height = Inches(1.85)


def _png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as handle:
        handle.seek(16)
        width, height = struct.unpack(">II", handle.read(8))
    return width, height


def _fit_box(
    pixel_w: int,
    pixel_h: int,
    *,
    max_width: float,
    max_height: float,
) -> tuple[float, float]:
    if pixel_w <= 0 or pixel_h <= 0:
        return max_width, max_height
    ratio = pixel_w / pixel_h
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def _layout_body(shape, *, split: bool) -> None:
    shape.left = Inches(_TEXT_LEFT)
    shape.top = Inches(_TEXT_TOP)
    shape.width = Inches(_TEXT_WIDTH_SPLIT if split else _TEXT_WIDTH_FULL)
    shape.height = Inches(_TEXT_HEIGHT)


def _add_image_fitted(slide, path: Path) -> bool:
    if not path.is_file():
        return False
    pixel_w, pixel_h = _png_size(path)
    width, height = _fit_box(
        pixel_w,
        pixel_h,
        max_width=_IMAGE_MAX_WIDTH,
        max_height=_IMAGE_MAX_HEIGHT,
    )
    left = _IMAGE_LEFT + (_IMAGE_MAX_WIDTH - width) / 2
    top = _IMAGE_TOP + (_IMAGE_MAX_HEIGHT - height) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    return True


def _presentation_story(meta: KkpMeta) -> list[tuple[str, str]]:
    """Sequential defense narrative aligned with the report structure."""
    return [
        (
            meta.topic,
            (
                f"{meta.student_full}\n"
                f"група {meta.group}\n\n"
                f"Керівник:\n{meta.advisor_line}\n\n"
                f"{meta.city} — {meta.year}"
            ),
        ),
        (
            "1. Проблема",
            (
                "Generative AI (Stable Diffusion, DALL·E, Midjourney) створює "
                "зображення, які важко відрізнити від фото.\n\n"
                "Класичні підказки (EXIF, JPEG-артефакти) слабшають.\n\n"
                "Потрібна проста система: завантажив фото — отримав підпис "
                "«real» / «AI» і відсоток впевненості."
            ),
        ),
        (
            "2. Мета і завдання",
            (
                f"Мета: програмна система «{meta.topic}» — навчання, оцінка "
                "і Gradio-demo для перевірки зображень.\n\n"
                "Завдання:\n"
                "• hi-res датасет real / AI\n"
                "• дві CNN: ResNet18 і EfficientNet-B0\n"
                "• порівняння метрик на test set\n"
                "• відкритий код і відтворювані експерименти"
            ),
        ),
        (
            "3. Ідея рішення",
            (
                "Бінарна класифікація — як тест «справжнє чи згенероване».\n\n"
                "Transfer learning: беремо CNN, уже навчену на ImageNet, "
                "і донавчаємо на парах real/AI (як учень, що вміє бачити "
                "форми, вчиться відрізняти два типи зображень).\n\n"
                "Вхід: JPEG/PNG, 224×224 px. Вихід: клас + confidence."
            ),
        ),
        (
            "4. Дані",
            (
                "Датасет Parveshiiii/AI-vs-Real [Hugging Face]:\n"
                "512×512 px, різні сцени (люди, об'єкти, інтер'єри).\n\n"
                "Структура: train / val / test, підкаталоги REAL і FAKE.\n\n"
                "Завантаження: make download-hires-hf у data/ai_hires/"
            ),
        ),
        (
            "5. Пайплайн системи",
            (
                "Конвеєр Makefile:\n"
                "1) download-hires-hf\n"
                "2) train / train-efficientnet\n"
                "3) evaluate-all — 204 test фото\n"
                "4) compare\n"
                "5) demo для комісії\n\n"
                "Справа — схема (рис. 3.2 звіту). "
                "configs/*.yaml — однакові умови для обох моделей."
            ),
        ),
        (
            "6. Дві моделі",
            (
                "ResNet18 — перша, проста модель для порівняння; швидше "
                "навчається на CPU.\n\n"
                "EfficientNet-B0 — уважніше до дрібних деталей на hi-res; "
                "додали після помилок ResNet18 на test set.\n\n"
                "Обидві: готові ваги ImageNet, замінений останній шар на 2 класи."
            ),
        ),
        (
            "7. Протокол експерименту",
            (
                "10 «уроків» навчання (повних проходів по train set).\n"
                "Під час train — штучні спотворення фото; на test — лише resize.\n\n"
                "Перевірка: 204 фото, яких модель не бачила (102 real + 102 AI).\n"
                "Зберігаємо best.pth за найкращою val — не за останнім уроком."
            ),
        ),
        (
            "8. Результати на test set",
            (
                "ResNet18: 93,6 %, F1 0,938\n"
                "• 10 хибних тривог\n"
                "• 3 пропущені підробки\n\n"
                "EfficientNet-B0: 98,5 %, F1 0,986\n"
                "• 3 хибні тривоги, 0 пропущених\n\n"
                "Справа — порівняння метрик. "
                "Різниця ≈10 фото з 204."
            ),
        ),
        (
            "9. Чому модель помиляється",
            (
                "Grad-CAM: куди CNN «дивилась».\n\n"
                "test_real_0049.jpg — ResNet 98,4 % AI, "
                "хоча фото справжнє (інтер'єр).\n\n"
                "Справа — heat map з табл. 5.3 звіту. "
                "EfficientNet прибрав пропуски підробок, "
                "але хибні тривоги на складних real залишились."
            ),
        ),
        (
            "10. Gradio demo",
            (
                "Інтерактивна частина для захисту:\n"
                "• upload або випадкове фото\n"
                "• ResNet18 / EfficientNet-B0\n"
                "• verdict + confidence\n\n"
                "Справа — головний екран (рис. 4.1). "
                "Той самий inference, що на test set."
            ),
        ),
        (
            "11. Якість і відтворюваність",
            (
                "pytest — модульні та інтеграційні тести.\n"
                "CI: Ruff + pytest + Docker build.\n\n"
                "make verify-results — перерахунок metrics з checkpoint-ів "
                "і порівняння з docs/artifacts/*.json.\n\n"
                "Цифри в звіті = цифри в коді (tol 1e-4)."
            ),
        ),
        (
            "12. Висновки",
            (
                "EfficientNet-B0 кращий на test set датасету; Gradio demo "
                "готовий до захисту.\n\n"
                "Практична цінність: відкритий код, Makefile, YAML, "
                "metrics.json.\n\n"
                f"GitHub: {meta.github}\n\n"
                "Подальший розвиток: більше даних, ensemble, REST API."
            ),
        ),
    ]


def _resolve_image_path(
    filename: str,
    *,
    assets_dir: Path,
    charts_dir: Path,
    comparison_dir: Path,
) -> Path:
    if filename == "metrics_comparison.png":
        for candidate in (charts_dir / filename, comparison_dir / filename):
            if candidate.is_file():
                return candidate
        return charts_dir / filename
    return assets_dir / filename


def fill_presentation(
    prs: Presentation,
    meta: KkpMeta,
    *,
    comparison_dir: Path,
    assets_dir: Path,
    charts_dir: Path | None = None,
) -> None:
    slides_content = _presentation_story(meta)
    if len(slides_content) != len(prs.slides):
        msg = (
            f"Template has {len(prs.slides)} slides, story has {len(slides_content)} — "
            "check Шаблон_презентації_до_ККП"
        )
        raise ValueError(msg)

    charts = charts_dir or comparison_dir
    accent = RGBColor(0x1A, 0x56, 0x8E)

    for slide_idx, (slide, (title, body)) in enumerate(
        zip(prs.slides, slides_content, strict=True)
    ):
        has_image = slide_idx in IMAGE_SLIDES
        shapes = [shape for shape in slide.shapes if shape.has_text_frame]
        if slide_idx == 0 and len(shapes) >= 2:
            _layout_title_slide(shapes[0], shapes[1])
            _set_shape_text(
                shapes[0], title, font_size=22, bold=True, color=accent, alignment=PP_ALIGN.CENTER
            )
            _set_shape_text(shapes[1], body, font_size=16, alignment=PP_ALIGN.CENTER)
        elif shapes:
            _set_shape_text(shapes[0], title, font_size=24, bold=True, color=accent)
            if len(shapes) > 1:
                _layout_body(shapes[1], split=has_image)
                font_size = _BODY_FONT_SPLIT if has_image else _BODY_FONT_FULL
                _set_shape_text(shapes[1], body, font_size=font_size)

        if has_image:
            image_path = _resolve_image_path(
                IMAGE_SLIDES[slide_idx],
                assets_dir=assets_dir,
                charts_dir=charts,
                comparison_dir=comparison_dir,
            )
            _add_image_fitted(slide, image_path)
