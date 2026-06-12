#!/usr/bin/env python3
"""Generate KKP report (DOCX) and defense presentation (PPTX) from official templates."""

from __future__ import annotations

# ruff: noqa: E402
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from capture_demo_screenshots import capture_demo_screenshots
from docx import Document
from docx_toc_pages import resolve_document_page_count, resolve_toc_pages
from export_report_artifacts import export_report_artifacts
from generate_report_figures import generate_report_figures
from kkp_docx_fill import fill_kkp_document, fill_table_of_contents, update_abstract_stats
from kkp_meta import DEFAULT_META
from kkp_pptx_fill import fill_presentation
from ooxml_utils import ppt_template_to_pptx, word_template_to_docx
from pptx import Presentation

DOCS = ROOT / "docs"
REPORT_ASSETS = DOCS / "report_assets"
COMPARISON = ROOT / "outputs" / "comparison"

WORD_TEMPLATE = ROOT / "Шаблон_записки_до_ККП_бакалавра_2025_2.dotm"
PPT_TEMPLATE = ROOT / "Шаблон_презентації_до_ККП_бакалавра_2025 (1).potx"

REPORT_NAME = "2026_Б_ККП_ПЗПІ-23-5_Ткач_Л_Я.docx"
PRESENTATION_NAME = "2026_Б_ККП_ПЗПІ-23-5_Ткач_Л_Я.pptx"


def build_report(output: Path) -> Path:
    tmp = DOCS / "_tmp_report.docx"
    work = DOCS / "_tmp_toc_work.docx"
    word_template_to_docx(WORD_TEMPLATE, tmp)
    doc = Document(str(tmp))
    fill_kkp_document(doc, DEFAULT_META, comparison_dir=COMPARISON, assets_dir=REPORT_ASSETS)
    doc.save(str(work))
    tmp.unlink(missing_ok=True)

    print("Resolving TOC page numbers from rendered layout...")
    toc_pages = resolve_toc_pages(work)
    doc = Document(str(work))
    fill_table_of_contents(doc, pages=toc_pages)
    doc.save(str(output))

    print("Updating abstract stats from rendered page count...")
    page_count = resolve_document_page_count(output)
    doc = Document(str(output))
    update_abstract_stats(doc, DEFAULT_META, pages=page_count)
    doc.save(str(output))
    print(f"Abstract stats: {page_count} pages")
    work.unlink(missing_ok=True)
    print(f"TOC page numbers: {toc_pages[0]}–{toc_pages[-1]} ({len(toc_pages)} entries)")

    rw_metrics = DOCS / "_tmp_real_world_metrics.json"
    try:
        from kkp.config import load_config
        from kkp.evaluate import evaluate_model, save_metrics

        config = load_config(ROOT / "configs" / "ai_generated_efficientnet.yaml")
        ckpt = ROOT / "outputs" / "ai_generated_efficientnet" / "checkpoints" / "best.pth"
        rw_dir = ROOT / "data" / "real_world"
        if ckpt.is_file() and rw_dir.is_dir():
            metrics, _ = evaluate_model(config, ckpt, data_dir=rw_dir)
            save_metrics(metrics, rw_metrics)
            export_report_artifacts(real_world_metrics=rw_metrics)
        else:
            export_report_artifacts()
    except FileNotFoundError as exc:
        print(f"Warning: {exc}")
        export_report_artifacts()
    finally:
        rw_metrics.unlink(missing_ok=True)

    print(f"Report artifacts: {DOCS / 'artifacts'}")
    return output


def build_presentation(output: Path) -> Path:
    tmp = DOCS / "_tmp_presentation.pptx"
    ppt_template_to_pptx(PPT_TEMPLATE, tmp)
    prs = Presentation(str(tmp))
    fill_presentation(
        prs,
        DEFAULT_META,
        comparison_dir=COMPARISON,
        assets_dir=REPORT_ASSETS,
        charts_dir=DOCS / "artifacts",
    )
    prs.save(str(output))
    tmp.unlink(missing_ok=True)
    return output


def main() -> None:
    DOCS.mkdir(exist_ok=True)
    report_path = DOCS / REPORT_NAME
    presentation_path = DOCS / PRESENTATION_NAME

    print("Generating report figures...")
    generate_report_figures(REPORT_ASSETS, data_root=ROOT / "data")
    print("Generating Grad-CAM figure...")
    try:
        from generate_gradcam import generate_gradcam_figure

        generate_gradcam_figure(
            REPORT_ASSETS / "gradcam_examples.png",
            data_root=ROOT / "data",
        )
    except Exception as exc:
        print(f"Warning: Grad-CAM figure skipped: {exc}")
    print("Capturing web UI screenshots from http://127.0.0.1:7860/ ...")
    try:
        capture_demo_screenshots(REPORT_ASSETS)
    except RuntimeError as exc:
        print(f"Warning: {exc}")
        print("Start the app (make demo) and run: make demo-screenshots")

    print("Building report from Word template...")
    build_report(report_path)
    print(f"Report: {report_path}")

    print("Building presentation from PPT template...")
    build_presentation(presentation_path)
    print(f"Presentation: {presentation_path}")


if __name__ == "__main__":
    main()
